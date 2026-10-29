"""Helper baca/tulis Excel dengan kolom yang harus diperlakukan sebagai teks."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def _existing_columns(path: Path) -> set[str]:
    header = pd.read_excel(path, nrows=0)
    return {str(col) for col in header.columns}


def read_excel(path: str | Path, text_columns: tuple[str, ...] = ()) -> pd.DataFrame:
    """Baca Excel dan paksa kolom tertentu bertipe string.

    Kolom teks dibaca dengan ``dtype=str`` supaya angka nol di depan
    (mis. nomor HP ``0812...``) tidak berubah menjadi angka.
    """
    path = Path(path)
    existing = _existing_columns(path)
    dtype = {col: str for col in text_columns if col in existing}
    return pd.read_excel(path, dtype=dtype or None)


def write_excel(
    df: pd.DataFrame,
    path: str | Path,
    text_columns: tuple[str, ...] = (),
) -> None:
    """Tulis DataFrame ke Excel, tandai kolom teks dengan number format ``@``."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(path, index=False)

    if not text_columns:
        return

    workbook = load_workbook(path)
    sheet = workbook.active
    header = {cell.value: cell.column for cell in sheet[1]}
    for column in text_columns:
        column_index = header.get(column)
        if column_index is None:
            continue
        for row in sheet.iter_rows(
            min_row=2, min_col=column_index, max_col=column_index
        ):
            for cell in row:
                cell.number_format = "@"
    workbook.save(path)

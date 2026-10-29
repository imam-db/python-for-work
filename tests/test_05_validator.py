"""Test untuk utility 05-data-validator."""

from pathlib import Path

import data_validator
import pandas as pd
from openpyxl import load_workbook

from common.config import load_config
from common.excel import read_excel

ROOT = Path(__file__).resolve().parents[1]


def load_rules() -> dict:
    return load_config(ROOT / "05-data-validator" / "config.yaml")["rules"]


def test_phone_column_is_read_as_text(tmp_path, load_utility):
    generate_sample = load_utility("05-data-validator", "generate_sample.py")
    path = generate_sample.generate_sample_data(tmp_path / "sample.xlsx")

    df = read_excel(path, text_columns=("No HP",))

    assert df["No HP"].iloc[0] == "081234567890"
    assert df["No HP"].dtype == object or str(df["No HP"].dtype) == "str"


def test_validator_finds_known_errors(tmp_path, load_utility):
    generate_sample = load_utility("05-data-validator", "generate_sample.py")
    path = generate_sample.generate_sample_data(tmp_path / "sample.xlsx")

    errors = data_validator.validate_data(str(path), load_rules())

    assert len(errors) == 17
    assert len({error["row"] for error in errors}) == 7

    phone_errors = [error for error in errors if error["rule"] == "phone"]
    assert len(phone_errors) == 2
    assert {error["value"] for error in phone_errors} == {"08123456", "0812"}


def test_phone_max_digits_is_enforced():
    df = pd.DataFrame({"No HP": ["081234567890123456789"]})
    validator = data_validator.Validator(df)

    errors = validator.validate({"No HP": [{"type": "phone"}]})

    assert len(errors) == 1
    assert "lebih dari 15 digit" in errors[0]["message"]


def test_report_keeps_leading_zero(tmp_path, load_utility):
    generate_sample = load_utility("05-data-validator", "generate_sample.py")
    path = generate_sample.generate_sample_data(tmp_path / "sample.xlsx")
    report = tmp_path / "report.xlsx"

    data_validator.validate_data(str(path), load_rules(), str(report))

    sheet = load_workbook(report)["Data"]
    header = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(row=2, column=header["No HP"]).value == "081234567890"
    assert sheet.cell(row=2, column=header["No HP"]).number_format == "@"

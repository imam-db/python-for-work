"""Tests for pfw.validate."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from pfw.common.config import load_config
from pfw.common.io import read_table
from pfw.sample import generate_validate
from pfw.validate import Validator, validate_data

ROOT = Path(__file__).resolve().parents[1]


def load_rules() -> dict:
    return load_config(ROOT / "configs" / "validate.yaml")["rules"]


def test_phone_column_is_read_as_text(tmp_path):
    path = generate_validate(tmp_path / "sample.xlsx")

    df = read_table(path, text_columns=("No HP",))

    assert df["No HP"].iloc[0] == "081234567890"
    assert df["No HP"].dtype == object or str(df["No HP"].dtype) == "str"


def test_validator_finds_known_errors(tmp_path):
    path = generate_validate(tmp_path / "sample.xlsx")

    errors = validate_data(str(path), load_rules())

    assert len(errors) == 17
    assert len({error["row"] for error in errors}) == 7

    phone_errors = [error for error in errors if error["rule"] == "phone"]
    assert len(phone_errors) == 2
    assert {error["value"] for error in phone_errors} == {"08123456", "0812"}


def test_phone_max_digits_is_enforced():
    df = pd.DataFrame({"No HP": ["081234567890123456789"]})
    validator = Validator(df)

    errors = validator.validate({"No HP": [{"type": "phone"}]})

    assert len(errors) == 1
    assert "lebih dari 15 digit" in errors[0]["message"]


def test_report_keeps_leading_zero(tmp_path):
    path = generate_validate(tmp_path / "sample.xlsx")
    report = tmp_path / "report.xlsx"

    validate_data(str(path), load_rules(), str(report))

    sheet = load_workbook(report)["Data"]
    header = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(row=2, column=header["No HP"]).value == "081234567890"
    assert sheet.cell(row=2, column=header["No HP"]).number_format == "@"

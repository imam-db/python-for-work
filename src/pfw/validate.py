"""Validate tabular rows against rules defined in a YAML config."""

from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pfw.common.config import load_config
from pfw.common.io import read_table
from pfw.common.logging_setup import get_logger, setup_logging
from pfw.common.phone import phone_digits

log = get_logger("validate")

DEFAULT_CONFIG = Path("configs") / "validate.yaml"

FILL_ERROR = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
FILL_HEADER = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")


class Validator:
    """Row validators driven by YAML rules."""

    def __init__(self, df: pd.DataFrame):
        self.df = df
        self.errors: list[dict] = []

    def validate_required(self, column: str, rule: dict) -> None:
        """Require a non-empty field."""
        for idx, value in self.df[column].items():
            if pd.isna(value) or str(value).strip() == "":
                self.errors.append(
                    {
                        "row": idx + 2,
                        "column": column,
                        "value": value,
                        "rule": "required",
                        "message": f"{column} must not be empty",
                    }
                )

    def validate_email(self, column: str, rule: dict) -> None:
        """Validate email format."""
        email_pattern = r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$"
        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                if not re.match(email_pattern, str(value)):
                    self.errors.append(
                        {
                            "row": idx + 2,
                            "column": column,
                            "value": value,
                            "rule": "email",
                            "message": f"Invalid email format: {value}",
                        }
                    )

    def validate_phone(self, column: str, rule: dict) -> None:
        """Validate phone numbers against min/max digit counts."""
        min_digits = rule.get("min_digits", 10)
        max_digits = rule.get("max_digits", 15)

        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                digits = phone_digits(value) or ""
                if len(digits) < min_digits:
                    message = f"No HP kurang dari {min_digits} digit: {value}"
                elif len(digits) > max_digits:
                    message = f"No HP lebih dari {max_digits} digit: {value}"
                else:
                    continue

                self.errors.append(
                    {
                        "row": idx + 2,
                        "column": column,
                        "value": value,
                        "rule": "phone",
                        "message": message,
                    }
                )

    def validate_date_range(self, column: str, rule: dict) -> None:
        """Validate dates against an optional range."""
        min_date = rule.get("min")
        max_date = rule.get("max")

        if min_date:
            min_date = datetime.strptime(min_date, "%Y-%m-%d")
        if max_date:
            max_date = datetime.strptime(max_date, "%Y-%m-%d")

        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                try:
                    if isinstance(value, datetime):
                        date_val = value
                    else:
                        date_val = pd.to_datetime(value)

                    if min_date and date_val < min_date:
                        self.errors.append(
                            {
                                "row": idx + 2,
                                "column": column,
                                "value": value,
                                "rule": "date_range",
                                "message": (
                                    f'Date before {min_date.strftime("%Y-%m-%d")}: {value}'
                                ),
                            }
                        )
                    elif max_date and date_val > max_date:
                        self.errors.append(
                            {
                                "row": idx + 2,
                                "column": column,
                                "value": value,
                                "rule": "date_range",
                                "message": (
                                    f'Date after {max_date.strftime("%Y-%m-%d")}: {value}'
                                ),
                            }
                        )
                except (ValueError, TypeError, OverflowError):
                    self.errors.append(
                        {
                            "row": idx + 2,
                            "column": column,
                            "value": value,
                            "rule": "date_range",
                            "message": f"Invalid date format: {value}",
                        }
                    )

    def validate_number_range(self, column: str, rule: dict) -> None:
        """Validate numbers against an optional range."""
        min_val = rule.get("min")
        max_val = rule.get("max")

        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                try:
                    num_val = float(value)
                    if min_val is not None and num_val < min_val:
                        self.errors.append(
                            {
                                "row": idx + 2,
                                "column": column,
                                "value": value,
                                "rule": "number_range",
                                "message": f"Nilai kurang dari {min_val}: {value}",
                            }
                        )
                    elif max_val is not None and num_val > max_val:
                        self.errors.append(
                            {
                                "row": idx + 2,
                                "column": column,
                                "value": value,
                                "rule": "number_range",
                                "message": f"Nilai lebih dari {max_val}: {value}",
                            }
                        )
                except (ValueError, TypeError):
                    self.errors.append(
                        {
                            "row": idx + 2,
                            "column": column,
                            "value": value,
                            "rule": "number_range",
                            "message": f"Bukan angka valid: {value}",
                        }
                    )

    def validate_regex(self, column: str, rule: dict) -> None:
        """Validate with a regex pattern."""
        pattern = rule.get("pattern")
        message = rule.get("message", f"Does not match format: {pattern}")

        if not pattern:
            return

        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                if not re.match(pattern, str(value)):
                    self.errors.append(
                        {
                            "row": idx + 2,
                            "column": column,
                            "value": value,
                            "rule": "regex",
                            "message": message,
                        }
                    )

    def validate_in_list(self, column: str, rule: dict) -> None:
        """Require the value to be one of an allowed list."""
        valid_values = rule.get("values", [])

        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                if str(value) not in [str(v) for v in valid_values]:
                    self.errors.append(
                        {
                            "row": idx + 2,
                            "column": column,
                            "value": value,
                            "rule": "in_list",
                            "message": (
                                f'Value "{value}" is not valid. '
                                f"Must be one of: {valid_values}"
                            ),
                        }
                    )

    def validate_unique(self, column: str, rule: dict) -> None:
        """Require unique values (no duplicates)."""
        seen = {}
        for idx, value in self.df[column].items():
            if pd.notna(value) and str(value).strip() != "":
                str_val = str(value)
                if str_val in seen:
                    self.errors.append(
                        {
                            "row": idx + 2,
                            "column": column,
                            "value": value,
                            "rule": "unique",
                            "message": f"Duplicate of row {seen[str_val]}: {value}",
                        }
                    )
                else:
                    seen[str_val] = idx + 2

    def validate(self, rules: dict) -> list:
        """Run every validation rule and collect errors."""
        for column, column_rules in rules.items():
            if column not in self.df.columns:
                log.warning("  Warning: column '%s' not found, skipping validation", column)
                continue

            for rule in column_rules:
                rule_type = rule.get("type")

                if rule_type == "required":
                    self.validate_required(column, rule)
                elif rule_type == "email":
                    self.validate_email(column, rule)
                elif rule_type == "phone":
                    self.validate_phone(column, rule)
                elif rule_type == "date_range":
                    self.validate_date_range(column, rule)
                elif rule_type == "number_range":
                    self.validate_number_range(column, rule)
                elif rule_type == "regex":
                    self.validate_regex(column, rule)
                elif rule_type == "in_list":
                    self.validate_in_list(column, rule)
                elif rule_type == "unique":
                    self.validate_unique(column, rule)

        return self.errors


def print_report(total_rows: int, errors: list) -> None:
    """Print the validation report to the console."""
    error_rows = set(e["row"] for e in errors)
    valid_count = total_rows - len(error_rows)

    log.info("")
    log.info("=" * 50)
    log.info("VALIDATION REPORT")
    log.info("=" * 50)

    log.info("")
    log.info("Summary:")
    log.info("  Total rows  : %s", total_rows)
    log.info("  Valid       : %s", valid_count)
    log.info("  Errors      : %s row(s) (%s issue(s))", len(error_rows), len(errors))

    if not errors:
        log.info("")
        log.info("All data is VALID!")
        return

    log.info("")
    log.info("--- ERROR DETAIL (%s) ---", len(errors))
    for error in errors[:30]:
        log.info("  Row %s: %s", error["row"], error["message"])

    if len(errors) > 30:
        log.info("  ... and %s more error(s)", len(errors) - 30)


def export_to_excel(
    df: pd.DataFrame, errors: list, output_path: str, text_columns: tuple[str, ...] = ()
) -> None:
    """Export the validation result to an .xlsx report (all values as text)."""
    workbook = Workbook()

    ws_summary = workbook.active
    ws_summary.title = "Summary"

    error_rows = set(e["row"] for e in errors)
    valid_count = len(df) - len(error_rows)

    ws_summary["A1"] = "Data Validation Report"
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary["A3"] = "Total rows:"
    ws_summary["B3"] = len(df)
    ws_summary["A4"] = "Valid:"
    ws_summary["B4"] = valid_count
    ws_summary["A5"] = "Error:"
    ws_summary["B5"] = len(error_rows)
    ws_summary["A6"] = "Total Masalah:"
    ws_summary["B6"] = len(errors)

    if errors:
        ws_errors = workbook.create_sheet("Detail Error")
        ws_errors.append(["Row", "Column", "Value", "Rule", "Message"])

        for cell in ws_errors[1]:
            cell.fill = FILL_HEADER
            cell.font = Font(bold=True, color="FFFFFF")

        for error in errors:
            ws_errors.append(
                [
                    error["row"],
                    error["column"],
                    str(error["value"]) if error["value"] is not None else "",
                    error["rule"],
                    error["message"],
                ]
            )

    ws_data = workbook.create_sheet("Data")
    ws_data.append(list(df.columns))
    for cell in ws_data[1]:
        cell.fill = FILL_HEADER
        cell.font = Font(bold=True, color="FFFFFF")

    error_cells = {(error["row"], error["column"]) for error in errors}
    text_indexes = [
        index for index, col in enumerate(df.columns, 1) if col in text_columns
    ]

    for idx, row in df.iterrows():
        excel_row = idx + 2
        row_data = [str(row[col]) if pd.notna(row[col]) else "" for col in df.columns]
        ws_data.append(row_data)

        for col_idx in text_indexes:
            ws_data.cell(row=excel_row, column=col_idx).number_format = "@"

        for col_idx, col in enumerate(df.columns, 1):
            if (excel_row, col) in error_cells:
                ws_data.cell(row=excel_row, column=col_idx).fill = FILL_ERROR
                ws_data.cell(row=excel_row, column=col_idx).number_format = "@"

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    log.info("")
    log.info("Report saved to: %s", output_path)


def phone_columns(rules: dict) -> tuple[str, ...]:
    """Return columns that have a ``phone`` rule."""
    columns = []
    for column, column_rules in rules.items():
        for rule in column_rules:
            if rule.get("type") == "phone":
                columns.append(column)
                break
    return tuple(columns)


def validate_data(input_file: str, rules: dict, output_file: str | None = None) -> list:
    """Validate a tabular file against YAML rules."""
    if not Path(input_file).exists():
        raise FileNotFoundError(f"File '{input_file}' not found")

    log.info("Reading file: %s", input_file)
    text_columns = phone_columns(rules)
    df = read_table(input_file, text_columns=text_columns)
    log.info("Total rows: %s", len(df))
    log.info("Columns: %s", ", ".join(map(str, df.columns)))

    log.info("")
    log.info("Validating data...")
    validator = Validator(df)
    errors = validator.validate(rules)

    print_report(len(df), errors)

    if output_file:
        export_to_excel(df, errors, output_file, text_columns)

    return errors


def add_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("input", nargs="?", help="Input tabular file")
    parser.add_argument("--output", "-o", help="Export the report to an .xlsx file")
    parser.add_argument(
        "--config",
        "-c",
        help="YAML config file (default: configs/validate.yaml)",
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="pfw validate",
        description="Validate tabular data against rules in a YAML config.",
    )
    add_arguments(parser)
    return parser


def run(args: argparse.Namespace) -> int:
    setup_logging()
    config_path = Path(args.config) if args.config else DEFAULT_CONFIG
    config = load_config(config_path)
    rules = config.get("rules", {})

    if args.input:
        input_file = Path(args.input)
    elif config.get("input"):
        input_file = Path(config["input"])
    else:
        input_file = None

    if args.output:
        output_file = Path(args.output)
    elif config.get("output"):
        output_file = Path(config["output"])
    else:
        output_file = None

    if not input_file:
        log.error("Error: input file is required (via CLI or config)")
        return 1

    if not rules:
        log.error("Error: validation rules must be defined in %s", config_path)
        return 1

    try:
        validate_data(
            str(input_file),
            rules,
            str(output_file) if output_file else None,
        )
    except (FileNotFoundError, ValueError) as exc:
        log.error("Error: %s", exc)
        return 1
    return 0


def main(argv: list[str] | None = None) -> int:
    return run(build_parser().parse_args(argv))

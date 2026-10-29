"""Read and write tabular files by extension with text-safe columns."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SUPPORTED_SUFFIXES = (".xlsx", ".csv", ".parquet")


def _check_suffix(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(SUPPORTED_SUFFIXES)
        raise ValueError(
            f"Format file '{suffix or path.name}' tidak didukung. Gunakan: {supported}"
        )
    return suffix


def read_table(
    path: str | Path,
    text_columns: tuple[str, ...] = (),
) -> pd.DataFrame:
    """Read an xlsx, csv, or parquet file and force text columns to string.

    Text columns keep leading zeros (for example phone numbers ``0812...``).
    """
    path = Path(path)
    suffix = _check_suffix(path)

    if suffix == ".xlsx":
        header = pd.read_excel(path, nrows=0)
        existing = {str(column) for column in header.columns}
        dtype = {column: str for column in text_columns if column in existing}
        return pd.read_excel(path, dtype=dtype or None)

    if suffix == ".csv":
        dtype = {column: str for column in text_columns}
        return pd.read_csv(path, dtype=dtype or None)

    frame = pd.read_parquet(path)
    for column in text_columns:
        if column in frame.columns:
            frame[column] = frame[column].astype("string")
    return frame


def write_table(
    df: pd.DataFrame,
    path: str | Path,
    text_columns: tuple[str, ...] = (),
) -> None:
    """Write a DataFrame to xlsx, csv, or parquet based on the file suffix.

    For xlsx, text columns are marked with the ``@`` number format so Excel
    keeps values such as ``081234567890`` as text.
    """
    path = Path(path)
    suffix = _check_suffix(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if suffix == ".xlsx":
        df.to_excel(path, index=False)
        _mark_text_columns(path, text_columns)
    elif suffix == ".csv":
        df.to_csv(path, index=False)
    else:
        frame = df.copy()
        for column in text_columns:
            if column in frame.columns:
                frame[column] = frame[column].astype("string")
        frame.to_parquet(path, index=False)


def _mark_text_columns(path: Path, text_columns: tuple[str, ...]) -> None:
    if not text_columns:
        return

    workbook = load_workbook(path)
    sheet = workbook.active
    header = {cell.value: cell.column for cell in sheet[1]}
    for column in text_columns:
        column_index = header.get(column)
        if column_index is None:
            continue
        for row in sheet.iter_rows(
            min_row=2, min_col=column_index, max_col=column_index
        ):
            for cell in row:
                cell.number_format = "@"
    workbook.save(path)

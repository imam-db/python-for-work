"""Compare two tabular files and report what changed."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pfw.common.io import read_table
from pfw.common.logging_setup import get_logger, setup_logging

log = get_logger("diff")


# Warna untuk highlight
FILL_ADDED = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")  # Hijau
FILL_DELETED = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")  # Merah
FILL_CHANGED = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")  # Kuning


def sort_key(value: str) -> tuple[int, int, str]:
    """Urutkan key numerik secara numerik, sisanya secara alfabetis."""
    text = str(value)
    try:
        return (0, int(text), "")
    except ValueError:
        return (1, 0, text)


def warn_duplicate_keys(df: pd.DataFrame, key_column: str, label: str) -> None:
    """Warn when the key column contains duplicate values."""
    keys = df[key_column].astype(str)
    duplicates = keys[keys.duplicated()].unique().tolist()
    if not duplicates:
        return

    preview = ", ".join(duplicates[:10])
    remaining = f" (+{len(duplicates) - 10} lainnya)" if len(duplicates) > 10 else ""
    log.warning(
        "Warning: %s duplicate key(s) in column '%s': %s%s; using the first row",
        label,
        key_column,
        preview,
        remaining,
    )


def compare_dataframes(
    df_old: pd.DataFrame, df_new: pd.DataFrame, key_column: str | None = None
) -> dict:
    """Bandingkan dua DataFrame dan return perbedaannya.

    If the key has duplicates, the first row is used for matching.
    """
    result: dict = {"added": [], "deleted": [], "changed": [], "summary": {}}

    if key_column and key_column in df_old.columns and key_column in df_new.columns:
        warn_duplicate_keys(df_old, key_column, "old file")
        warn_duplicate_keys(df_new, key_column, "new file")

        old_keys = set(df_old[key_column].astype(str))
        new_keys = set(df_new[key_column].astype(str))

        # Added rows (in new, not in old)
        for key in sorted(new_keys - old_keys, key=sort_key):
            row = df_new[df_new[key_column].astype(str) == key].iloc[0].to_dict()
            result["added"].append({"key": key, "data": row})

        # Deleted rows (in old, not in new)
        for key in sorted(old_keys - new_keys, key=sort_key):
            row = df_old[df_old[key_column].astype(str) == key].iloc[0].to_dict()
            result["deleted"].append({"key": key, "data": row})

        # Shared keys: check cell-level changes
        for key in sorted(old_keys & new_keys, key=sort_key):
            old_row = df_old[df_old[key_column].astype(str) == key].iloc[0]
            new_row = df_new[df_new[key_column].astype(str) == key].iloc[0]

            for col in df_old.columns:
                if col not in df_new.columns:
                    continue
                old_val = old_row[col]
                new_val = new_row[col]

                if pd.isna(old_val) and pd.isna(new_val):
                    continue
                if str(old_val) != str(new_val):
                    result["changed"].append(
                        {
                            "key": key,
                            "column": col,
                            "old_value": old_val,
                            "new_value": new_val,
                        }
                    )
    else:
        max_rows = max(len(df_old), len(df_new))

        for i in range(max_rows):
            if i >= len(df_old):
                result["added"].append({"row": i + 1, "data": df_new.iloc[i].to_dict()})
            elif i >= len(df_new):
                result["deleted"].append({"row": i + 1, "data": df_old.iloc[i].to_dict()})
            else:
                for col in df_old.columns:
                    if col not in df_new.columns:
                        continue
                    old_val = df_old.iloc[i][col]
                    new_val = df_new.iloc[i][col]

                    if pd.isna(old_val) and pd.isna(new_val):
                        continue
                    if str(old_val) != str(new_val):
                        result["changed"].append(
                            {
                                "row": i + 1,
                                "column": col,
                                "old_value": old_val,
                                "new_value": new_val,
                            }
                        )

    result["summary"] = {
        "added_count": len(result["added"]),
        "deleted_count": len(result["deleted"]),
        "changed_count": len(result["changed"]),
    }

    return result


def print_report(result: dict, key_column: str | None = None) -> None:
    """Print the comparison report to the console."""
    summary = result["summary"]

    log.info("")
    log.info("=" * 50)
    log.info("COMPARISON REPORT")
    log.info("=" * 50)

    log.info("")
    log.info("Summary:")
    log.info("  - Rows added    : %s", summary["added_count"])
    log.info("  - Rows deleted  : %s", summary["deleted_count"])
    log.info("  - Cells changed : %s", summary["changed_count"])

    if not (summary["added_count"] or summary["deleted_count"] or summary["changed_count"]):
        log.info("")
        log.info("Both files are IDENTICAL; no differences.")
        return

    if result["changed"]:
        log.info("")
        log.info("--- CHANGES (%s) ---", len(result["changed"]))
        for item in result["changed"][:20]:
            if key_column:
                log.info(
                    "  [%s=%s] %s: %s -> %s",
                    key_column,
                    item["key"],
                    item["column"],
                    item["old_value"],
                    item["new_value"],
                )
            else:
                log.info(
                    "  Row %s, %s: %s -> %s",
                    item["row"],
                    item["column"],
                    item["old_value"],
                    item["new_value"],
                )
        if len(result["changed"]) > 20:
            log.info("  ... and %s more change(s)", len(result["changed"]) - 20)

    if result["added"]:
        log.info("")
        log.info("--- ADDED ROWS (%s) ---", len(result["added"]))
        for item in result["added"][:10]:
            if key_column:
                log.info("  [%s=%s]", key_column, item["key"])
            else:
                log.info("  Row %s", item["row"])
        if len(result["added"]) > 10:
            log.info("  ... and %s more row(s)", len(result["added"]) - 10)

    if result["deleted"]:
        log.info("")
        log.info("--- DELETED ROWS (%s) ---", len(result["deleted"]))
        for item in result["deleted"][:10]:
            if key_column:
                log.info("  [%s=%s]", key_column, item["key"])
            else:
                log.info("  Row %s", item["row"])
        if len(result["deleted"]) > 10:
            log.info("  ... and %s more row(s)", len(result["deleted"]) - 10)


def export_to_excel(
    result: dict, output_path: str, key_column: str | None = None
) -> None:
    """Export the comparison result to a highlighted spreadsheet report."""
    workbook = Workbook()

    ws_summary = workbook.active
    ws_summary.title = "Summary"
    ws_summary["A1"] = "Comparison Report"
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary["A3"] = "Rows added:"
    ws_summary["B3"] = result["summary"]["added_count"]
    ws_summary["A4"] = "Rows deleted:"
    ws_summary["B4"] = result["summary"]["deleted_count"]
    ws_summary["A5"] = "Cells changed:"
    ws_summary["B5"] = result["summary"]["changed_count"]

    if result["changed"]:
        ws_changed = workbook.create_sheet("Changes")
        if key_column:
            ws_changed.append([key_column, "Column", "Old value", "New value"])
        else:
            ws_changed.append(["Row", "Column", "Old value", "New value"])

        for item in result["changed"]:
            if key_column:
                ws_changed.append(
                    [item["key"], item["column"], str(item["old_value"]), str(item["new_value"])]
                )
            else:
                ws_changed.append(
                    [item["row"], item["column"], str(item["old_value"]), str(item["new_value"])]
                )

        for cell in ws_changed[1]:
            cell.fill = FILL_CHANGED
            cell.font = Font(bold=True)

    if result["added"]:
        ws_added = workbook.create_sheet("New Rows")
        headers = list(result["added"][0]["data"].keys())
        ws_added.append(headers)
        for cell in ws_added[1]:
            cell.fill = FILL_ADDED
            cell.font = Font(bold=True)

        for item in result["added"]:
            ws_added.append(list(item["data"].values()))

    if result["deleted"]:
        ws_deleted = workbook.create_sheet("Deleted Rows")
        headers = list(result["deleted"][0]["data"].keys())
        ws_deleted.append(headers)
        for cell in ws_deleted[1]:
            cell.fill = FILL_DELETED
            cell.font = Font(bold=True)

        for item in result["deleted"]:
            ws_deleted.append(list(item["data"].values()))

    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    log.info("")
    log.info("Report saved to: %s", output_path)


def compare_excel(
    file_old: str,
    file_new: str,
    key_column: str | None = None,
    output_file: str | None = None,
) -> dict:
    """Compare two tabular files and optionally export a report."""
    if not Path(file_old).exists():
        raise FileNotFoundError(f"File '{file_old}' not found")
    if not Path(file_new).exists():
        raise FileNotFoundError(f"File '{file_new}' not found")

    log.info("Old file : %s", file_old)
    log.info("New file : %s", file_new)
    if key_column:
        log.info("Key column: %s", key_column)

    df_old = read_table(file_old)
    df_new = read_table(file_new)

    log.info("")
    log.info("Old file: %s rows, %s columns", len(df_old), len(df_old.columns))
    log.info("New file: %s rows, %s columns", len(df_new), len(df_new.columns))

    if key_column:
        if key_column not in df_old.columns:
            raise ValueError(
                f"Column '{key_column}' not found in the old file. "
                f"Available columns: {', '.join(map(str, df_old.columns))}"
            )
        if key_column not in df_new.columns:
            raise ValueError(f"Column '{key_column}' not found in the new file")

    result = compare_dataframes(df_old, df_new, key_column)
    print_report(result, key_column)

    if output_file:
        export_to_excel(result, output_file, key_column)

    return result


def add_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("file_old", help="Older tabular file (.xlsx, .csv, .parquet)")
    parser.add_argument("file_new", help="Newer tabular file (.xlsx, .csv, .parquet)")
    parser.add_argument("--key", "-k", help="Key column used to match rows")
    parser.add_argument("--output", "-o", help="Export the report to an .xlsx file")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="pfw diff",
        description="Compare two tabular files and show the differences.",
    )
    add_arguments(parser)
    return parser


def run(args: argparse.Namespace) -> int:
    setup_logging()
    try:
        compare_excel(
            args.file_old,
            args.file_new,
            args.key,
            args.output,
        )
    except (FileNotFoundError, ValueError) as exc:
        log.error("Error: %s", exc)
        return 1
    return 0


def main(argv: list[str] | None = None) -> int:
    return run(build_parser().parse_args(argv))
